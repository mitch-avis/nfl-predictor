"""Tests for the Excel betting template generator."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

from nfl_predictor.reporting.betting_excel import write_betting_template_xlsx


def test_write_betting_template_xlsx_creates_workbook(tmp_path: Path) -> None:
    """Writes a workbook with expected sheets and formulas."""
    # Intentionally unsorted input; template should sort by kickoff datetime, then game_id.
    df = pd.DataFrame(
        [
            {
                "season": 2025,
                "week": 19,
                "date": "2026-01-12",
                "gametime": "20:00",
                "game_datetime": "2026-01-12T20:00:00",
                "game_id": "2025_19_X_Y",
                "away_abbr": "X",
                "home_abbr": "Y",
                "home_win_prob": 0.62,
                "predicted_margin": 3.5,
                "predicted_total": 45.0,
                "home_moneyline": -135,
                "away_moneyline": 115,
                "home_spread": -2.5,
                "away_spread": 2.5,
                "total_line": 44.5,
            },
            {
                "season": 2025,
                "week": 19,
                "date": "2026-01-11",
                "gametime": "20:00",
                "game_datetime": "2026-01-11T20:00:00",
                "game_id": "2025_19_M_N",
                "away_abbr": "M",
                "home_abbr": "N",
                "home_win_prob": 0.53,
                "predicted_margin": 1.0,
                "predicted_total": 41.0,
                "home_moneyline": -110,
                "away_moneyline": -110,
                "home_spread": -1.0,
                "away_spread": 1.0,
                "total_line": 41.5,
            },
            {
                "season": 2025,
                "week": 19,
                "date": "2026-01-11",
                "gametime": "13:00",
                "game_datetime": "2026-01-11T13:00:00",
                "game_id": "2025_19_A_B",
                "away_abbr": "A",
                "home_abbr": "B",
                "home_win_prob": 0.41,
                "predicted_margin": -2.0,
                "predicted_total": 51.0,
                "home_moneyline": 120,
                "away_moneyline": -140,
                "home_spread": 3.0,
                "away_spread": -3.0,
                "total_line": 50.5,
            },
        ]
    )

    out = tmp_path / "template.xlsx"
    write_betting_template_xlsx(predictions=df, out_path=out)
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    assert "README" in wb.sheetnames
    assert "Bets" in wb.sheetnames
    assert "Live" in wb.sheetnames

    ws = wb["Bets"]
    headers = [c.value for c in ws[1]]

    assert headers[:8] == [
        "season",
        "week",
        "date",
        "gametime",
        "game_datetime",
        "game_id",
        "away_abbr",
        "home_abbr",
    ]

    # Row 2 should include a moneyline->prob formula.
    col_idx = headers.index("market_home_prob_raw") + 1
    cell = ws.cell(row=2, column=col_idx)
    assert isinstance(cell.value, str)
    assert cell.value.startswith("=IF(")

    # Spread/total odds should default to -110 so downstream formulas don't start as N/A.
    assert ws.cell(row=2, column=headers.index("home_spread_odds_live") + 1).value == -110
    assert ws.cell(row=2, column=headers.index("away_spread_odds_live") + 1).value == -110
    assert ws.cell(row=2, column=headers.index("total_over_odds_live") + 1).value == -110
    assert ws.cell(row=2, column=headers.index("total_under_odds_live") + 1).value == -110

    # Key live-input columns should exist for usability.
    for col in (
        "away_spread_live",
        "away_spread_odds_live",
        "home_spread_live",
        "home_spread_odds_live",
        "away_moneyline_live",
        "home_moneyline_live",
        "total_live",
        "total_over_odds_live",
        "total_under_odds_live",
        "money_action",
        "spread_action",
        "total_action",
    ):
        assert col in headers

    # Sorted order: earliest kickoff first (A_B at 13:00, then M_N at 20:00), then 2026-01-12.
    game_id_col = headers.index("game_id") + 1
    assert ws.cell(row=2, column=game_id_col).value == "2025_19_A_B"
    assert ws.cell(row=3, column=game_id_col).value == "2025_19_M_N"
    assert ws.cell(row=4, column=game_id_col).value == "2025_19_X_Y"

    # Live sheet should include a win-probability formula cell.
    ws_live = wb["Live"]
    live_headers = [c.value for c in ws_live[1]]
    assert "live_home_win_prob" in live_headers
    assert "total_over_odds_live" in live_headers

    # Live identifiers should reference the matching Bets rows (no matchup misalignment).
    live_game_id_col = live_headers.index("game_id") + 1
    live_away_abbr_col = live_headers.index("away_abbr") + 1
    live_home_abbr_col = live_headers.index("home_abbr") + 1
    assert ws_live.cell(row=2, column=live_game_id_col).value == "='Bets'!F2"
    assert ws_live.cell(row=2, column=live_away_abbr_col).value == "='Bets'!G2"
    assert ws_live.cell(row=2, column=live_home_abbr_col).value == "='Bets'!H2"
    assert ws_live.cell(row=3, column=live_game_id_col).value == "='Bets'!F3"
    assert ws_live.cell(row=3, column=live_away_abbr_col).value == "='Bets'!G3"
    assert ws_live.cell(row=3, column=live_home_abbr_col).value == "='Bets'!H3"

    # Live user-input defaults for fast in-game updates.
    assert ws_live.cell(row=2, column=live_headers.index("quarter") + 1).value == 1
    assert (
        ws_live.cell(row=2, column=live_headers.index("minutes_remaining_in_quarter") + 1).value
        == 15
    )
    assert ws_live.cell(row=2, column=live_headers.index("away_score_live") + 1).value == 0
    assert ws_live.cell(row=2, column=live_headers.index("home_score_live") + 1).value == 0

    # Live odds defaults.
    assert ws_live.cell(row=2, column=live_headers.index("home_spread_odds_live") + 1).value == -110
    assert ws_live.cell(row=2, column=live_headers.index("away_spread_odds_live") + 1).value == -110
    assert ws_live.cell(row=2, column=live_headers.index("total_over_odds_live") + 1).value == -110
    assert ws_live.cell(row=2, column=live_headers.index("total_under_odds_live") + 1).value == -110

    # Live moneylines should initialize from the dataset/model inputs (not -110).
    assert ws_live.cell(row=2, column=live_headers.index("away_moneyline_live") + 1).value == -140
    assert ws_live.cell(row=2, column=live_headers.index("home_moneyline_live") + 1).value == 120
    live_prob_cell = ws_live.cell(row=2, column=live_headers.index("live_home_win_prob") + 1)
    assert isinstance(live_prob_cell.value, str)
    assert "NORMSDIST" in live_prob_cell.value

    # Date should be present for sorting.
    date_cell = ws.cell(row=2, column=headers.index("date") + 1)
    assert str(date_cell.value) == "2026-01-11"

    kickoff_cell = ws.cell(row=2, column=headers.index("game_datetime") + 1)
    assert str(kickoff_cell.value).startswith("2026-01-11")

    # Autofit should write explicit widths for visible columns.
    assert ws.column_dimensions["A"].width is not None

    # Helper columns should remain hidden.
    helper_idx = headers.index("spread_mu_margin") + 1
    helper_letter = get_column_letter(helper_idx)
    assert ws.column_dimensions[helper_letter].hidden is True


def _sample_predictions() -> pd.DataFrame:
    """Return a two-game prediction frame with the columns the template consumes."""
    return pd.DataFrame(
        [
            {
                "season": 2025,
                "week": 19,
                "date": "2026-01-11",
                "gametime": "20:00",
                "game_datetime": "2026-01-11T20:00:00",
                "game_id": "2025_19_M_N",
                "away_abbr": "M",
                "home_abbr": "N",
                "home_win_prob": 0.55,
                "predicted_margin": 1.5,
                "predicted_total": 43.0,
                "home_moneyline": -120,
                "away_moneyline": 100,
                "home_spread": -1.5,
                "away_spread": 1.5,
                "total_line": 43.5,
            },
            {
                "season": 2025,
                "week": 19,
                "date": "2026-01-12",
                "gametime": "20:00",
                "game_datetime": "2026-01-12T20:00:00",
                "game_id": "2025_19_X_Y",
                "away_abbr": "X",
                "home_abbr": "Y",
                "home_win_prob": 0.62,
                "predicted_margin": 3.5,
                "predicted_total": 45.0,
                "home_moneyline": -135,
                "away_moneyline": 115,
                "home_spread": -2.5,
                "away_spread": 2.5,
                "total_line": 44.5,
            },
        ]
    )


def test_every_written_formula_parses_as_valid_excel(tmp_path: Path) -> None:
    """Excel silently discards a malformed formula, so none may be emitted.

    A formula whose parentheses do not nest makes Excel report the workbook as
    corrupt and strip the offending cells during repair, leaving those columns blank
    with no error raised anywhere in this pipeline. Tokenizing rather than counting
    characters also catches nesting that balances overall but closes too early.
    """
    out = tmp_path / "betting.xlsx"
    write_betting_template_xlsx(predictions=_sample_predictions(), out_path=out)

    wb = openpyxl.load_workbook(out)
    offenders: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                depth = 0
                for token in Tokenizer(value).items:
                    if token.subtype == "OPEN":
                        depth += 1
                    elif token.subtype == "CLOSE":
                        depth -= 1
                    if depth < 0:
                        offenders.append(f"{sheet_name}!{cell.coordinate} closes too early")
                        break
                else:
                    if depth != 0:
                        offenders.append(f"{sheet_name}!{cell.coordinate} unclosed ({depth})")

    assert not offenders, "malformed formulas: " + "; ".join(offenders[:5])


def test_action_columns_are_populated_for_every_game(tmp_path: Path) -> None:
    """The action columns are the point of the report, so they must carry formulas."""
    out = tmp_path / "betting.xlsx"
    write_betting_template_xlsx(predictions=_sample_predictions(), out_path=out)

    wb = openpyxl.load_workbook(out)
    expected = {
        "Bets": ("spread_action", "money_action", "total_action"),
        "Live": ("live_spread_action", "live_money_action", "live_total_action"),
    }
    for sheet_name, columns in expected.items():
        ws = wb[sheet_name]
        header = {str(cell.value): index for index, cell in enumerate(ws[1], start=1) if cell.value}
        for column in columns:
            assert column in header, f"{sheet_name} is missing {column}"
            for excel_row in range(2, ws.max_row + 1):
                value = ws.cell(row=excel_row, column=header[column]).value
                assert isinstance(value, str) and value.startswith("="), (
                    f"{sheet_name}!{column} row {excel_row} has no formula"
                )


def test_action_formulas_cover_every_recommendation_band(tmp_path: Path) -> None:
    """Each action formula names all five bands, so no edge falls through unlabeled."""
    out = tmp_path / "betting.xlsx"
    write_betting_template_xlsx(predictions=_sample_predictions(), out_path=out)

    wb = openpyxl.load_workbook(out)
    ws = wb["Bets"]
    header = {str(cell.value): index for index, cell in enumerate(ws[1], start=1) if cell.value}
    formula = ws.cell(row=2, column=header["spread_action"]).value
    assert isinstance(formula, str)

    for band in ("PASS", "LEAN", "SMALL", "MEDIUM", "STRONG"):
        assert band in formula, f"{band} missing from the action formula"
